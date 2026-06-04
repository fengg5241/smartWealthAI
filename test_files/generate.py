"""Generate test PDF and DOCX files for multi-tenant RAG demo."""

from pathlib import Path

OUT_DIR = Path(__file__).parent

ACME_CONTENT = """ACME Corp 产品手册

一、公司概览

ACME Corp（简称 ACME）成立于 1998 年，总部位于上海市浦东新区张江高科技园区，是一家专注于高端精密零部件制造的企业。公司董事长兼首席执行官为张明远。现任首席技术官为李芳华博士，首席财务官为陈志强。2025 年全年营收达到 48.6 亿元人民币，员工总数超过 3200 人。

二、产品线与价格

产品一：ACME-P100 精密轴承。规格为内径 10mm、外径 30mm、精度等级 P4。单价为人民币 185 元每套，最小起订量 100 套，交货周期 15 个工作日。

产品二：ACME-S200 不锈钢弹簧。规格为线径 0.5 至 3.0mm、材质 304 或 316 不锈钢。单价为人民币 8.5 元每件，最小起订量 500 件，交货周期 7 个工作日。

产品三：ACME-G300 精密齿轮。规格为模数 0.5 至 2.0、精度等级 ISO 5 级。单价为人民币 420 元每套，最小起订量 50 套，交货周期 20 个工作日。

产品四：ACME-H400 液压阀体。规格为工作压力 0 至 31.5MPa、材质球墨铸铁。单价为人民币 1280 元每件，最小起订量 10 件，交货周期 30 个工作日。

三、质量标准与认证

ACME Corp 全线产品通过 ISO 9001:2015 质量管理体系认证。汽车零部件产品通过 IATF 16949:2016 认证。出口产品符合 RoHS 和 REACH 环保标准。2023 年获得上海市专精特新企业称号。不良品率控制在 50 PPM 以下。

四、保修与售后服务政策

标准保修期为自发货之日起 24 个月。精密轴承 ACME-P100 系列享受延长保修至 36 个月。保修期内出现非人为质量问题，免费更换新品。售后技术支持热线为 021-5888-1234。紧急技术支持 7 乘 24 小时手机为 138-0000-5678。

五、联系方式

总部地址为上海市浦东新区张江路 1688 号 ACME 大厦 12 层，邮编 201203。销售咨询邮箱为 sales@acme-corp.cn，技术咨询邮箱为 tech@acme-corp.cn，官网为 www.acme-corp.cn，传真为 021-5888-1235。

六、重要合作伙伴

ACME Corp 长期为上海大众汽车有限公司、西门子医疗系统有限公司、博世中国投资有限公司、霍尼韦尔中国有限公司等企业提供精密零部件供应服务。
"""

GLOBEX_CONTENT = """Globex Inc 员工手册（2025 版）

第一章 公司简介与文化

Globex Inc 是一家全球领先的企业级 SaaS 软件公司，总部位于美国加利福尼亚州圣何塞市。中国区总部设在深圳市南山区科技园南区，成立于 2015 年。公司中国区总裁为王思远，技术副总裁为赵晓东博士。Globex 的使命是用技术连接商业，核心价值观为客户第一、创新驱动、多元包容。

第二章 考勤与工时制度

标准工作时间为周一至周五上午九点至下午六点，午休十二点至下午一点。实行弹性工作制，可申请上午八点至十点到岗。每月允许远程办公不超过八天。迟到三十分钟以内不计入考勤异常，但每月累计不超过三次。加班需提前在 OA 系统提交申请并获得直属上级审批。周末及法定节假日加班享受双倍工资或调休。

第三章 薪酬与福利

薪酬结构为基本工资加绩效奖金加年终奖，年终奖为十三至十五薪。五险一金按实际工资全额缴纳，非最低基数。补充商业医疗保险覆盖员工本人及配偶和子女。年度体检标准为 1500 元每人。员工持股计划 ESOP 覆盖所有正式员工，入职满一年即可参与。每月交通补贴 800 元，通讯补贴 300 元。午餐补贴每日 50 元，加班晚餐报销上限 80 元每餐。

第四章 年假与假期政策

法定年假 5 天，入职满 1 年后每年增加 1 天，上限 15 天。带薪病假每年 12 天。婚假 10 个工作日，需提供结婚证明。女性员工产假 158 天，男性员工陪产假 15 天。每年提供 3 天自我提升假，可用于参加行业会议或培训。入职满 3 年可申请为期 1 个月的创新休假，用于个人项目探索。

第五章 安全与应急规范

办公区域全面禁烟，包括电子烟。消防疏散集合点为大厦一楼南广场。每季度组织一次消防演练，全体员工必须参加。急救药箱位于各楼层茶水间入口处。禁止携带易燃易爆物品进入办公区域。紧急情况拨打内线 999 联系安保部。

第六章 IT 设备使用规定

公司配发的笔记本电脑为固定资产，离职时需归还。禁止在办公设备上安装未授权软件或游戏。公司内网 WiFi 仅限认证设备接入，SSID 为 Globex-Secure。敏感文件须加密存储，禁止通过个人邮箱或即时通讯工具传输。VPN 登录使用双因素认证 2FA，密码每 90 天更换一次。技术部门员工可申请额外显示器上限 2 台和开发服务器资源。

第七章 联系方式

深圳总部地址为深圳市南山区科技园南路 88 号环球科技大厦 25 层，邮编 518057。HR 咨询邮箱为 hr-china@globex.com。IT 支持热线为 0755-3688-1000 工作时间或 itsupport@globex.com。员工投诉与建议匿名邮箱为 feedback@globex.com。公司官网为 www.globex.com。

本手册最终解释权归 Globex Inc 中国区人力资源部所有，如有修订以最新版本为准。修订日期为 2025 年 03 月 01 日。
"""


def generate_pdf():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("SimHei", fname=r"C:\Windows\Fonts\simhei.ttf")
    pdf.set_font("SimHei", size=10)

    for line in ACME_CONTENT.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(4)
            continue
        if line.startswith(("一、", "二、", "三、", "四、", "五、", "六、")):
            pdf.ln(2)
            pdf.set_font("SimHei", size=12)
            pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("SimHei", size=10)
        else:
            pdf.multi_cell(0, 5.5, line)

    out = OUT_DIR / "acme_product_manual.pdf"
    pdf.output(str(out))
    print(f"Generated: {out} ({round(Path(out).stat().st_size/1024, 1)} KB)")


def generate_docx():
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(10.5)

    for line in GLOBEX_CONTENT.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph("")
            continue
        p = doc.add_paragraph()
        if line.startswith(("第一章", "第二章", "第三章", "第四章", "第五章", "第六章", "第七章")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(14)
        elif line.startswith("Globex Inc") and "手册" in line:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(16)
        else:
            p.add_run(line)

    out = OUT_DIR / "globex_handbook.docx"
    doc.save(str(out))
    print(f"Generated: {out} ({round(Path(out).stat().st_size/1024, 1)} KB)")


ACME_CONTENT_EN = """ACME Corp Product Catalog

Section 1 - Company Overview

ACME Corp (ACME) was founded in 1998 and is headquartered in Zhangjiang Hi-Tech Park, Pudong New Area, Shanghai, China. The company specializes in high-precision mechanical component manufacturing. The Chairman and CEO is Zhang Mingyuan, the CTO is Dr. Li Fanghua, and the CFO is Chen Zhiqiang. In 2025, annual revenue reached 4.86 billion RMB, with over 3,200 employees.

Section 2 - Products and Pricing

Product A: ACME-P100 Precision Bearing. Specifications include an inner diameter of 10mm, outer diameter of 30mm, and precision grade P4. Unit price is 185 RMB per set, minimum order quantity is 100 sets, and delivery time is 15 working days.

Product B: ACME-S200 Stainless Steel Spring. Specifications include wire diameter 0.5 to 3.0mm, material 304 or 316 stainless steel. Unit price is 8.5 RMB per piece, minimum order quantity is 500 pieces, and delivery time is 7 working days.

Product C: ACME-G300 Precision Gear. Specifications include module 0.5 to 2.0 and precision grade ISO Class 5. Unit price is 420 RMB per set, minimum order quantity is 50 sets, and delivery time is 20 working days.

Product D: ACME-H400 Hydraulic Valve Body. Specifications include working pressure 0 to 31.5MPa and ductile iron material. Unit price is 1,280 RMB per unit, minimum order quantity is 10 units, and delivery time is 30 working days.

Section 3 - Quality Standards and Certifications

All ACME Corp products have passed ISO 9001:2015 Quality Management System certification. Automotive components have passed IATF 16949:2016 certification. Export products comply with RoHS and REACH environmental standards. In 2023, the company was awarded the Shanghai Specialized and New Enterprise title. The defect rate is controlled under 50 PPM.

Section 4 - Warranty and After-Sales Service

The standard warranty period is 24 months from the date of shipment. The ACME-P100 Precision Bearing series enjoys an extended warranty of 36 months. Non-artificial quality issues during the warranty period qualify for free replacement. The after-sales support hotline is 021-5888-1234. Emergency technical support is available 7 by 24 at mobile 138-0000-5678.

Section 5 - Contact Information

Headquarters address is 12F, ACME Tower, 1688 Zhangjiang Road, Pudong New Area, Shanghai, China, zip code 201203. Sales inquiry email is sales@acme-corp.com, technical inquiry email is tech@acme-corp.com, website is www.acme-corp.com, fax is 021-5888-1235.

Section 6 - Key Partners

ACME Corp has long-term partnerships with SAIC Volkswagen Automotive Co Ltd, Siemens Healthineers Ltd, Bosch China Investment Ltd, and Honeywell China Ltd for precision component supply.
"""

GLOBEX_CONTENT_EN = """Globex Inc Employee Handbook (2025 Edition)

Chapter 1 - Company Overview

Globex Inc is a leading global enterprise SaaS software company headquartered in San Jose, California, USA. The China office is located in Nanshan Science Park, Shenzhen, and was established in 2015. The China President is Wang Siyuan, and the VP of Technology is Dr. Zhao Xiaodong. Globex's mission is to connect business with technology, and its core values are customer first, innovation-driven, and diversity and inclusion.

Chapter 2 - Attendance and Working Hours

Standard working hours are Monday to Friday from 9 AM to 6 PM, with lunch break from 12 PM to 1 PM. A flexible working system allows employees to arrive between 8 AM and 10 AM. Remote work is permitted for up to 8 days per month. Being late by less than 30 minutes is not counted as an attendance violation, but no more than 3 times per month. Overtime requires prior approval through the OA system from the direct supervisor. Weekend and public holiday overtime qualifies for double pay or compensatory time off.

Chapter 3 - Compensation and Benefits

The compensation structure consists of base salary plus performance bonus plus year-end bonus (13 to 15 months). Social insurance and housing fund are paid at the full actual salary, not the minimum base. Supplemental commercial medical insurance covers the employee, spouse, and children. The annual physical examination standard is 1,500 RMB per person. The Employee Stock Ownership Plan (ESOP) covers all regular employees, available after one year of employment. Monthly transportation subsidy is 800 RMB, and communication subsidy is 300 RMB. Daily lunch subsidy is 50 RMB, and overtime dinner reimbursement is capped at 80 RMB per meal.

Chapter 4 - Leave and Holiday Policy

Statutory annual leave is 5 days, increasing by 1 day each year after the first year of employment, up to a maximum of 15 days. Paid sick leave is 12 days per year. Marriage leave is 10 working days with a marriage certificate. Female employees receive 158 days of maternity leave, and male employees receive 15 days of paternity leave. Each year, 3 days of self-improvement leave are provided for attending industry conferences or training. Employees with 3 or more years of service can apply for a 1-month innovation leave for personal project exploration.

Chapter 5 - Safety and Emergency Procedures

The office area is completely smoke-free, including e-cigarettes. The fire evacuation assembly point is the south plaza on the first floor of the building. Fire drills are organized quarterly, and all employees must participate. First aid kits are located at the pantry entrance on each floor. Flammable and explosive items are prohibited in the office area. In emergencies, dial extension 999 to contact the Security Department.

Chapter 6 - IT Equipment Usage Policy

Company-issued laptops are fixed assets and must be returned upon resignation. Installing unauthorized software or games on company equipment is prohibited. The internal Wi-Fi is only accessible by authorized devices, with the SSID Globex-Secure. Sensitive files must be encrypted and must not be transmitted via personal email or instant messaging tools. VPN login requires two-factor authentication (2FA), and passwords must be changed every 90 days. Technical department employees may apply for additional monitors (up to 2) and development server resources.

Chapter 7 - Contact Information

Shenzhen headquarters address is 25F, Global Technology Tower, 88 Keyuan South Road, Nanshan District, Shenzhen, China, zip code 518057. HR inquiry email is hr-china@globex.com. IT support hotline is 0755-3688-1000 during work hours or itsupport@globex.com. Employee complaints and suggestions anonymous email is feedback@globex.com. Company website is www.globex.com.

This handbook is for internal reference only. The final interpretation right belongs to Globex Inc China HR Department. Any revisions will be based on the latest version. Revision date is March 1, 2025.
"""


def generate_pdf_en():
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font("SimHei", fname=r"C:\Windows\Fonts\simhei.ttf")
    pdf.set_font("SimHei", size=10)

    for line in ACME_CONTENT_EN.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(4)
            continue
        if line.startswith(("Section", "Chapter")):
            pdf.ln(2)
            pdf.set_font("SimHei", size=12)
            pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("SimHei", size=10)
        else:
            pdf.multi_cell(0, 5.5, line)

    out = OUT_DIR / "acme_product_manual_en.pdf"
    pdf.output(str(out))
    print(f"Generated: {out} ({round(Path(out).stat().st_size/1024, 1)} KB)")


def generate_docx_en():
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10.5)

    for line in GLOBEX_CONTENT_EN.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph("")
            continue
        p = doc.add_paragraph()
        if line.startswith(("Chapter", "Globex Inc")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(14) if line.startswith("Chapter") else Pt(16)
        else:
            p.add_run(line)

    out = OUT_DIR / "globex_handbook_en.docx"
    doc.save(str(out))
    print(f"Generated: {out} ({round(Path(out).stat().st_size/1024, 1)} KB)")


if __name__ == "__main__":
    generate_pdf()
    generate_docx()
    generate_pdf_en()
    generate_docx_en()
